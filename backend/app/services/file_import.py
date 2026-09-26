import csv
import io
import re
from zipfile import BadZipFile
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.core.config import Settings
from app.models.order import OrderSource, OrderStatus
from app.schemas.extraction import ExtractionOutcome
from app.schemas.order import OrderCreate, OrderItemCreate
from app.services.extraction import HybridOrderExtractionService, extraction_confidence, missing_required_fields


class FileImportError(Exception):
    def __init__(self, message: str, status_code: int = 422) -> None:
        super().__init__(message)
        self.message = message
        self.status_code = status_code


HEADER_ALIASES = {
    "sku": {"sku", "kod", "kod produktu", "product code"},
    "quantity": {"quantity", "qty", "ilosc", "ilość"},
    "unit": {"unit", "jednostka", "jm"},
    "product_name": {"product name", "product_name", "produkt", "nazwa produktu"},
    "customer_name": {"customer", "customer name", "customer_name", "klient", "nazwa klienta"},
    "customer_email": {"customer email", "customer_email", "email", "e-mail"},
    "customer_tax_id": {"customer tax id", "customer_tax_id", "tax id", "nip", "vat"},
    "delivery_address": {"delivery address", "delivery_address", "adres dostawy", "adres"},
    "delivery_date": {"delivery date", "delivery_date", "data dostawy"},
}


def normalize_header(value: Any) -> str:
    text = re.sub(r"[_\-]+", " ", str(value or "").strip().casefold())
    return re.sub(r"\s+", " ", text)


def canonical_header(value: Any) -> str | None:
    normalized = normalize_header(value)
    for canonical, aliases in HEADER_ALIASES.items():
        if normalized in aliases:
            return canonical
    return None


class SpreadsheetOrderExtractor:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings

    def extract(self, content: bytes, filename: str, source: OrderSource) -> ExtractionOutcome:
        if source == OrderSource.CSV:
            rows = self._csv_rows(content)
        elif source == OrderSource.XLSX:
            rows = self._xlsx_rows(content)
        else:
            raise FileImportError("Nieobsługiwany format arkusza", 415)
        return self._rows_to_order(rows, filename, source)

    def _csv_rows(self, content: bytes) -> list[dict[str, Any]]:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise FileImportError("Plik CSV musi być zapisany w UTF-8") from exc
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
            reader = csv.reader(io.StringIO(text), dialect)
            return self._tabular_rows(list(reader))
        except csv.Error as exc:
            raise FileImportError("Nie udało się odczytać pliku CSV") from exc

    def _xlsx_rows(self, content: bytes) -> list[dict[str, Any]]:
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            raw_rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            workbook.close()
        except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
            raise FileImportError("Nie udało się odczytać pliku XLSX") from exc
        return self._tabular_rows(raw_rows)

    def _tabular_rows(self, rows: list[list[Any]]) -> list[dict[str, Any]]:
        rows = [row for row in rows if any(value not in (None, "") for value in row)]
        if not rows:
            raise FileImportError("Plik nie zawiera danych")
        if len(rows) - 1 > self.settings.max_spreadsheet_rows:
            raise FileImportError(
                f"Plik może zawierać maksymalnie {self.settings.max_spreadsheet_rows} pozycji",
                413,
            )
        headers = [canonical_header(value) for value in rows[0]]
        if "sku" not in headers or "quantity" not in headers:
            raise FileImportError("Arkusz musi zawierać kolumny SKU i quantity")
        if len([header for header in headers if header]) != len(set(header for header in headers if header)):
            raise FileImportError("Arkusz zawiera zduplikowane kolumny")
        return [
            {header: row[index] if index < len(row) else None for index, header in enumerate(headers) if header}
            for row in rows[1:]
        ]

    def _rows_to_order(
        self,
        rows: list[dict[str, Any]],
        filename: str,
        source: OrderSource,
    ) -> ExtractionOutcome:
        warnings: list[str] = []
        items: list[OrderItemCreate] = []
        for row_number, row in enumerate(rows, start=2):
            sku = self._text(row.get("sku"))
            quantity = self._quantity(row.get("quantity"))
            if not sku or quantity is None:
                warnings.append(f"Pominięto wiersz {row_number}: wymagane są SKU i dodatnia ilość")
                continue
            try:
                items.append(
                    OrderItemCreate(
                        sku=sku,
                        product_name=self._text(row.get("product_name")),
                        quantity=quantity,
                        unit=self._text(row.get("unit")) or "szt.",
                    )
                )
            except ValidationError:
                warnings.append(f"Pominięto nieprawidłowy wiersz {row_number}")

        metadata: dict[str, Any] = {}
        for field in (
            "customer_name",
            "customer_email",
            "customer_tax_id",
            "delivery_address",
            "delivery_date",
        ):
            values = self._unique_values(rows, field)
            if len(values) > 1:
                warnings.append(f"Niejednoznaczne pole: {field}")
                metadata[field] = None
            elif values:
                metadata[field] = values[0]
            else:
                metadata[field] = None
        metadata["delivery_date"] = self._date(metadata["delivery_date"], warnings)

        try:
            order = OrderCreate(
                **metadata,
                source=source,
                status=OrderStatus.NEEDS_REVIEW,
                original_content=f"Imported file: {Path(filename).name}",
                items=items,
            )
        except ValidationError as exc:
            raise FileImportError("Dane klienta w pliku mają nieprawidłowy format") from exc
        missing = missing_required_fields(order)
        warnings.extend(f"Brak wymaganego pola: {field}" for field in missing)
        order.status = OrderStatus.NEW if not missing and not warnings else OrderStatus.NEEDS_REVIEW
        return ExtractionOutcome(
            order=order,
            warnings=warnings,
            confidence=extraction_confidence(order),
        )

    @staticmethod
    def _text(value: Any) -> str | None:
        if value is None:
            return None
        text = str(value).strip()
        return text or None

    @staticmethod
    def _quantity(value: Any) -> int | None:
        if isinstance(value, bool):
            return None
        try:
            number = float(str(value).replace(",", "."))
        except (TypeError, ValueError):
            return None
        return int(number) if number > 0 and number.is_integer() else None

    @staticmethod
    def _unique_values(rows: list[dict[str, Any]], field: str) -> list[Any]:
        values: list[Any] = []
        for row in rows:
            value = row.get(field)
            if value in (None, ""):
                continue
            normalized = value.strip() if isinstance(value, str) else value
            if normalized not in values:
                values.append(normalized)
        return values

    @staticmethod
    def _date(value: Any, warnings: list[str]) -> date | None:
        if value is None:
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = str(value).strip()
        for pattern in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
            try:
                return datetime.strptime(text, pattern).date()
            except ValueError:
                continue
        warnings.append("Nieprawidłowa data dostawy")
        return None


class PDFTextExtractor:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings

    def extract_text(self, content: bytes) -> str:
        try:
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted and reader.decrypt("") == 0:
                raise FileImportError("Zabezpieczony PDF nie może zostać odczytany")
            if len(reader.pages) > self.settings.max_pdf_pages:
                raise FileImportError(
                    f"PDF może zawierać maksymalnie {self.settings.max_pdf_pages} stron",
                    413,
                )
            text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
        except (PdfReadError, OSError, ValueError) as exc:
            raise FileImportError("Nie udało się odczytać pliku PDF") from exc
        if len(text) < 10:
            raise FileImportError(
                "PDF nie zawiera tekstu możliwego do odczytania. Skanowane dokumenty wymagają OCR."
            )
        return text


class FileImportService:
    def __init__(
        self,
        settings: Settings,
        spreadsheet_extractor: SpreadsheetOrderExtractor,
        pdf_extractor: PDFTextExtractor,
        text_extraction: HybridOrderExtractionService,
    ) -> None:
        self.settings = settings
        self.spreadsheet_extractor = spreadsheet_extractor
        self.pdf_extractor = pdf_extractor
        self.text_extraction = text_extraction

    def extract(self, content: bytes, filename: str) -> ExtractionOutcome:
        if not content:
            raise FileImportError("Przesłany plik jest pusty")
        if len(content) > self.settings.max_upload_size_bytes:
            limit_mb = self.settings.max_upload_size_bytes // (1024 * 1024)
            raise FileImportError(f"Plik przekracza limit {limit_mb} MB", 413)
        suffix = Path(filename).suffix.casefold()
        if suffix == ".csv":
            return self.spreadsheet_extractor.extract(content, filename, OrderSource.CSV)
        if suffix == ".xlsx":
            return self.spreadsheet_extractor.extract(content, filename, OrderSource.XLSX)
        if suffix == ".pdf":
            text = self.pdf_extractor.extract_text(content)
            return self.text_extraction.extract(text, OrderSource.PDF)
        raise FileImportError("Obsługiwane formaty plików: CSV, XLSX i PDF", 415)
